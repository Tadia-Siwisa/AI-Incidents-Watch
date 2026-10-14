#!/usr/bin/env python3
"""Build the evidence-aware Wide Search 2 discovery workbook.

The workbook is a discovery register, not a confirmed-incident count. It keeps
the source collections' own classifications, records overlaps with the existing
46-family baseline, and identifies records whose cited sources were unavailable
from the research environment.
"""

from __future__ import annotations

from collections import Counter
from copy import deepcopy
from datetime import datetime
import hashlib
import json
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "wide_search_2"
EVIDENCE = ROOT / "evidence" / "wide_search_2" / "2026-10-14"
OUTPUT = ROOT / "AI_Incident_Watch_WIDE_Search_2.xlsx"
REGISTER_OUTPUT = DATA / "wide_search_2_register_2026-10-14.json"

PERMISSION_FILE = DATA / "permission_protocol_records_2026-10-05.json"
METR_FILE = DATA / "metr_agent_incidents_2026-05-19.json"
MANUAL_FILE = DATA / "manual_cases_2026-10-14.json"
REACHABILITY_FILE = EVIDENCE / "permission_protocol_source_reachability.json"
SEARCH_LOG_FILE = DATA / "search_log_2026-10-14.json"

AS_OF = "2026-10-14"
PERMISSION_TRACKER = "https://permissionprotocol.com/agent-incident-tracker"
METR_TRACKER = "https://metr.org/agent-incidents/"
METR_REPORT = "https://metr.org/blog/2026-05-19-frontier-risk-report/"


# Exact family relationships found during reconciliation. Unlisted records are
# left as new candidates; fuzzy title similarity is never used as adjudication.
PP_RELATIONSHIPS = {
    "openai-agents-wikimedia-unauthorized-edits-etherpad": (
        "Existing dated backfill", "BF-001"
    ),
    "openai-agents-53-chatgpt-user-images-external-hosting": (
        "Existing baseline family", "W06"
    ),
    "openai-agent-dns-sandbox-bypass": (
        "Existing baseline family", "N17"
    ),
    "google-gemini-security-agents-breached-three-firms": (
        "Existing baseline family", "N09"
    ),
    "australia-gym-api-idor-agentic-waitlist-manipulation-august2026": (
        "Existing baseline family", "N04"
    ),
    "meta-muse-spark-unauthorized-api-spams-august2026": (
        "Existing baseline family", "W21"
    ),
    "aisi-uncontrolled-supply-chain-intervention-august2026": (
        "Existing baseline family", "N15"
    ),
    "openai-irregular-eval-sandbox-breakout": (
        "Potential overlap; adjudication required", "W22"
    ),
    "anthropic-claude-breached-3-organizations-security-evals": (
        "Existing baseline umbrella", "N06; N07; N08"
    ),
    "hugging-face-autonomous-ai-agent-breach-july2026": (
        "Existing baseline family", "N14"
    ),
    "openai-agent-australian-medicare-portal-breach": (
        "Existing baseline family", "N12"
    ),
    "pocketos-production-database-backups-deleted": (
        "Existing baseline family", "N05"
    ),
    "replit-agent-production-data-deletion": (
        "Existing baseline family", "N01"
    ),
}


# Direct sources and corrections added after opening the underlying material.
PP_ENRICHMENTS = {
    "openai-agents-wikimedia-unauthorized-edits-etherpad": {
        "sources": [
            {
                "label": "Wikimedia Foundation — OpenAI agent activity (retrieved route)",
                "url": "https://diff.wikimedia.org/2026/10/05/openai-rogue-agent-activities-found-on-wikimedia-projects/",
            }
        ],
        "note": "Existing backfill BF-001 already preserves Wikimedia's finding that no system or data compromise was found and that outage causation was only possible, not established.",
    },
    "google-gemini-security-agents-breached-three-firms": {
        "sources": [
            {"label": "BBC — Gemini accessed three companies", "url": "https://www.bbc.com/news/articles/c607l0k72rlvo"}
        ],
        "note": "Existing family N09 preserves that the company identities and model versions were not disclosed.",
    },
    "australia-gym-api-idor-agentic-waitlist-manipulation-august2026": {
        "sources": [
            {"label": "BBC — gym booking account", "url": "https://www.bbc.com/news/articles/cn0nww2qlp7o"}
        ],
        "note": "Existing family N04 treats this as an attributed operational account without a published gym-side audit.",
    },
    "meta-muse-spark-unauthorized-api-spams-august2026": {
        "sources": [
            {"label": "Meta — third-party test misconfiguration", "url": "https://research.meta.ai/blog/addressing-third-party-testing-misconfiguration-muse-spark-1-1/"}
        ],
        "note": "Existing family W21 contains the retrieved Meta disclosure.",
    },
    "aisi-uncontrolled-supply-chain-intervention-august2026": {
        "sources": [
            {"label": "UK AISI incident report", "url": "https://www.aisi.gov.uk/blog/incident-report-unsanctioned-agent-behaviour-during-cyber-testing"}
        ],
        "note": "Existing family N15 records 19 unsuccessful attempts in an internet-enabled evaluation and AISI's statement that no real-world harm occurred.",
    },
    "anthropic-claude-breached-3-organizations-security-evals": {
        "sources": [
            {"label": "Anthropic — alignment-assessment cybersecurity incidents", "url": "https://www.anthropic.com/research/alignment-assessment-cybersecurity-incidents"}
        ],
        "note": "The Permission Protocol links for this row were unavailable. The existing N06–N08 families use Anthropic's retrieved incident report and should control final wording and family counts.",
    },
    "openai-agent-australian-medicare-portal-breach": {
        "sources": [
            {"label": "OpenAI — how we will do better for Australia", "url": "https://openai.com/index/how-we-will-do-better-for-australia/"}
        ],
        "note": "Existing family N12 distinguishes the Medicare statistics portal from core Medicare systems and does not infer patient-record access.",
    },
    "pocketos-production-database-backups-deleted": {
        "sources": [
            {"label": "The Register — PocketOS production deletion", "url": "https://www.theregister.com/software/2026/04/27/cursor-opus-agent-snuffs-out-startups-production-database/"}
        ],
        "note": "Existing family N05 contains the retrieved account and its qualifications.",
    },
    "replit-agent-production-data-deletion": {
        "sources": [
            {"label": "The Register — Replit production database deletion", "url": "https://www.theregister.com/software/2025/07/21/vibe-coding-service-replit-deleted-production-database/"}
        ],
        "note": "Existing family N01 treats this as a detailed attributed operational report because no Replit technical root-cause report was found.",
    },
    "datatalks-terraform-production-deletion": {
        "sources": [
            {"label": "Rootly interview with Alexey Grigorev", "url": "https://rootly.com/humans-of-reliability/hey-claude-wheres-my-database-alexey-grigorev"},
            {"label": "Tom's Hardware — DataTalks report", "url": "https://www.tomshardware.com/tech-industry/artificial-intelligence/claude-code-deletes-developers-production-setup-including-its-database-and-snapshots-2-5-years-of-records-were-nuked-in-an-instant"},
        ],
        "note": "Correction from the founder interview: he said he does not blame Claude, saw the terraform destroy command and did not stop it, and recovered the database after AWS restored an internal backup. Treat this as human-approved execution and process/configuration failure context.",
    },
    "amazon-kiro-aws-outage-report": {
        "sources": [
            {"label": "Amazon — correction and response", "url": "https://www.aboutamazon.com/news/aws/aws-service-outage-ai-bot-kiro"}
        ],
        "note": "Disputed. Amazon says the brief interruption resulted from a misconfigured role and user error, calls AI involvement coincidental, and denies a claimed second AWS event. The Financial Times account was blocked in this environment.",
    },
    "amazon-q-developer-vscode-system-prompt-wiper-july2025": {
        "sources": [
            {"label": "AWS Security Bulletin AWS-2025-015", "url": "https://aws.amazon.com/security/security-bulletins/AWS-2025-015/"},
            {"label": "GitHub advisory GHSA-7g7f-ff96-5gcw", "url": "https://github.com/aws/aws-toolkit-vscode/security/advisories/GHSA-7g7f-ff96-5gcw"},
        ],
        "note": "AWS says malicious code reached extension version 1.84.0 but failed to execute because of a syntax error, preventing changes to services or customer environments. This is a supply-chain incident and an attempted agent payload, not a completed AI deletion.",
    },
    "echoleak-cve-2025-32711-m365-copilot-zero-click-exfiltration": {
        "sources": [
            {"label": "EchoLeak research paper", "url": "https://arxiv.org/html/2509.10540v1"},
            {"label": "Microsoft CVE-2025-32711", "url": "https://msrc.microsoft.com/update-guide/vulnerability/CVE-2025-32711"},
        ],
        "note": "A working researcher proof of concept demonstrated the chain. The research says Microsoft found no evidence of exploitation in the wild; classify it as a demonstrated production-product vulnerability, not a confirmed customer data breach.",
    },
    "gemini-cli-cvss-10-rce-ghsa-wpqr-6v78-jr5g": {
        "sources": [
            {"label": "GitHub advisory GHSA-wpqr-6v78-jr5g", "url": "https://github.com/advisories/GHSA-wpqr-6v78-jr5g"}
        ],
        "note": "The advisory describes a critical path to remote code execution in untrusted CI workspaces. It does not report confirmed in-the-wild exploitation.",
    },
}


# URLs for which this project already contains a retrieved source or a source
# opened during the present pass. These augment the batch link-access check.
LOCALLY_VERIFIED_URLS = {
    "https://www.anthropic.com/research/alignment-assessment-cybersecurity-incidents",
    "https://www.aisi.gov.uk/blog/incident-report-unsanctioned-agent-behaviour-during-cyber-testing",
    "https://research.meta.ai/blog/addressing-third-party-testing-misconfiguration-muse-spark-1-1/",
    "https://www.bbc.com/news/articles/cn0nww2qlp7o",
    "https://www.bbc.com/news/articles/c607l0k72rlvo",
    "https://openai.com/index/how-we-will-do-better-for-australia/",
    "https://www.theregister.com/software/2026/04/27/cursor-opus-agent-snuffs-out-startups-production-database/",
    "https://www.theregister.com/software/2025/07/21/vibe-coding-service-replit-deleted-production-database/",
    "https://diff.wikimedia.org/2026/10/05/openai-rogue-agent-activities-found-on-wikimedia-projects/",
    "https://rootly.com/humans-of-reliability/hey-claude-wheres-my-database-alexey-grigorev",
    "https://www.tomshardware.com/tech-industry/artificial-intelligence/claude-code-deletes-developers-production-setup-including-its-database-and-snapshots-2-5-years-of-records-were-nuked-in-an-instant",
    "https://www.aboutamazon.com/news/aws/aws-service-outage-ai-bot-kiro",
    "https://aws.amazon.com/security/security-bulletins/AWS-2025-015/",
    "https://github.com/aws/aws-toolkit-vscode/security/advisories/GHSA-7g7f-ff96-5gcw",
    "https://arxiv.org/html/2509.10540v1",
    "https://msrc.microsoft.com/update-guide/vulnerability/CVE-2025-32711",
    "https://github.com/advisories/GHSA-wpqr-6v78-jr5g",
}


HEADERS = [
    "WS2 ID",
    "Date / reference date",
    "Date reported",
    "Case family / report",
    "Source collection",
    "Environment",
    "Actor / system",
    "What happened",
    "Outcome / impact",
    "Real external action?",
    "Boundary / scope dimension",
    "Evidence status",
    "Source quality",
    "Relationship to existing corpus",
    "Existing ID(s)",
    "Review status",
    "Limitations / corrections",
    "Primary source",
    "Additional sources",
    "Stable source record",
]


HEADER_FILL = PatternFill("solid", fgColor="17365D")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
REALIZED_FILL = PatternFill("solid", fgColor="FCE4D6")
CONTROLLED_FILL = PatternFill("solid", fgColor="E2F0D9")
SCREENED_FILL = PatternFill("solid", fgColor="E7E6E6")
OVERLAP_FILL = PatternFill("solid", fgColor="FFF2CC")
UNREACHABLE_FILL = PatternFill("solid", fgColor="F4CCCC")
THIN_GREY = Side(style="thin", color="D9E1F2")


def load_json(path: Path):
    return json.loads(path.read_text())


def unique_sources(sources: list[dict]) -> list[dict]:
    seen = set()
    result = []
    for source in sources:
        url = source.get("url", "").strip()
        if not url or url in seen:
            continue
        seen.add(url)
        result.append({"label": source.get("label") or url, "url": url})
    return result


def pp_rows(permission: dict, reachability: dict) -> list[dict]:
    status_by_url = {row["url"]: row for row in reachability["results"]}
    rows = []
    for index, original in enumerate(permission["records"], 1):
        item = deepcopy(original)
        enrichment = PP_ENRICHMENTS.get(item["id"], {})
        sources = unique_sources(item.get("sources", []) + enrichment.get("sources", []))
        reachable = []
        blocked = []
        for source in sources:
            check = status_by_url.get(source["url"])
            if source["url"] in LOCALLY_VERIFIED_URLS or (check and check["status"] == "reachable"):
                reachable.append(source["url"])
            elif check and check["status"] != "reachable":
                blocked.append(source["url"])
        relation, existing = PP_RELATIONSHIPS.get(
            item["id"], ("New candidate record", "")
        )
        correction = enrichment.get("note", "")
        limitations = " ".join(
            part for part in [item.get("reviewNote", "").strip(), correction.strip()] if part
        )
        if not limitations:
            limitations = (
                "No correction is recorded in the imported source register. "
                "Independent claim-level adjudication is still required."
            )
        if reachable:
            access = f"{len(reachable)} linked source(s) reachable; source-linked discovery, not independently re-adjudicated"
        else:
            access = "No linked supporting source was reachable in this environment; retain as an unverified lead"
        if relation.startswith("Existing"):
            review = "Existing adjudicated family; use the baseline/update finding for analysis"
        elif relation.startswith("Potential"):
            review = "Possible duplicate; compare source event and dates before counting"
        elif not reachable:
            review = "Hold — source unavailable"
        elif item["incidentKind"] == "demonstrated":
            review = "Controlled/security context — adjudication pending"
        else:
            review = "Reported event — adjudication pending"
        rows.append(
            {
                "ws2_id": f"PP-{index:03d}",
                "source_key": item["id"],
                "reference_date": item.get("date") or "Unknown",
                "reported_date": "Unknown — publisher's date may be occurrence or disclosure",
                "headline": item["headline"],
                "collection": "Permission Protocol",
                "environment": (
                    "Reported real-world event"
                    if item["incidentKind"] == "realized"
                    else "Controlled demonstration / vulnerability research"
                ),
                "actor": item.get("tool") or "Unknown",
                "what_happened": item.get("summary") or item.get("reportedAction") or "Unknown",
                "outcome": item.get("productionImpact") or "Unknown",
                "real_external_action": (
                    "Reported/realized; exact agent role varies"
                    if item["incidentKind"] == "realized"
                    else "No confirmed deployment incident; demonstrated"
                ),
                "boundary": item.get("category") or "Unknown",
                "evidence_status": access,
                "source_quality": item.get("sourceQuality") or "Unknown",
                "relationship": relation,
                "existing_ids": existing,
                "review_status": review,
                "limitations": limitations,
                "sources": sources,
                "stable_record": item.get("stableUrl") or PERMISSION_TRACKER,
                "source_access_count": len(reachable),
                "source_access_fail_count": len(blocked),
                "kind": item["incidentKind"],
                "severity": item.get("severity"),
                "publisher_fields": {
                    "operator": item.get("operator"),
                    "incidentType": item.get("incidentType"),
                    "productionSurface": item.get("productionSurface"),
                    "relevantGate": item.get("relevantGate"),
                    "missingAuthorityCheck": item.get("missingAuthorityCheck"),
                },
            }
        )
    return rows


def metr_rows(items: list[dict]) -> list[dict]:
    rows = []
    for item in items:
        if item["id"].startswith("INC-H"):
            continue
        number = int(item["id"].split("-")[1])
        if number <= 6:
            environment = "Screened context — error/fabrication without boundary action"
            real_action = "No boundary-crossing action identified"
            review = "Screened out of incident set; retained for transparency"
        elif 19 <= number <= 21:
            environment = "Internal operational deployment"
            real_action = "Internal tool/system action reported"
            review = "New source-linked family — independent adjudication pending"
        elif number in (38, 39):
            environment = "Attributed operational report from an unnamed company"
            real_action = "Reported operational effect"
            review = "New qualified family — company and trace undisclosed"
        else:
            environment = "Controlled evaluation"
            real_action = (
                "Controlled run with limited real external effects"
                if number == 43
                else "No confirmed deployment incident; controlled run"
            )
            review = "Controlled context — independent adjudication pending"
        summary = item.get("summary", {})
        source_url = item.get("source_url") or METR_REPORT
        source_label = item.get("source") or "METR Frontier Risk Report"
        reported = "2026-03-19" if 19 <= number <= 21 else "2026-05-19"
        action = summary.get("actions_against_intent") or "Unknown"
        task = summary.get("task") or "Task not disclosed."
        outcome = summary.get("outcomes_impacts") or "No separate outcome was reported."
        limitations = (
            "METR preserves a source quotation and scoring rationale. This pass verified the METR register and linked source route, "
            "but did not independently reproduce every run."
        )
        if number in (38, 39):
            limitations += " The participating company was not named."
        rows.append(
            {
                "ws2_id": "METR-" + item["id"].split("-")[1],
                "source_key": item["id"],
                "reference_date": "Unknown",
                "reported_date": reported,
                "headline": f"{item.get('model', 'Unknown model')}: {action[:180]}",
                "collection": "METR documented-agent-incident register",
                "environment": environment,
                "actor": item.get("model") or "Unknown",
                "what_happened": f"Task: {task} Reported action: {action}",
                "outcome": outcome,
                "real_external_action": real_action,
                "boundary": (item.get("overreach") or {}).get("bucket") or "No overreach bucket",
                "evidence_status": "Source-linked curated account with quoted passage; not independently re-adjudicated in this pass",
                "source_quality": item.get("source_short") or "METR curated source",
                "relationship": "New candidate record",
                "existing_ids": "",
                "review_status": review,
                "limitations": limitations,
                "sources": unique_sources(
                    [
                        {"label": source_label, "url": source_url},
                        {"label": "METR documented incidents register", "url": METR_TRACKER},
                    ]
                ),
                "stable_record": METR_TRACKER + "#" + item["id"].lower(),
                "source_access_count": 1,
                "source_access_fail_count": 0,
                "kind": "screened" if number <= 6 else ("operational" if number in (19, 20, 21, 38, 39) else "controlled"),
                "severity": None,
                "publisher_fields": {
                    "sourceQuote": item.get("source_quotes"),
                    "overreach": item.get("overreach"),
                    "deception": item.get("deception"),
                    "graded": item.get("graded"),
                },
            }
        )
    return rows


def manual_rows(payload: dict) -> list[dict]:
    rows = []
    for item in payload["cases"]:
        environment = item["environment"]
        if "Controlled" in environment or "demonstration" in environment.lower():
            kind = "controlled"
            real_action = "No confirmed unapproved deployment incident"
        elif "game" in environment.lower():
            kind = "authorized_game"
            real_action = "Real transfer within an authorized adversarial game"
        else:
            kind = "operational"
            real_action = "Reported real-system action; qualifications apply"
        rows.append(
            {
                "ws2_id": item["id"],
                "source_key": item["id"],
                "reference_date": item["activity_date"],
                "reported_date": item["reported_date"],
                "headline": item["headline"],
                "collection": "Direct expanded search",
                "environment": environment,
                "actor": item["actor"],
                "what_happened": item["what_happened"],
                "outcome": item["outcome"],
                "real_external_action": real_action,
                "boundary": item["boundary_dimension"],
                "evidence_status": item["evidence_status"],
                "source_quality": item["source_quality"],
                "relationship": item["relationship"],
                "existing_ids": item["existing_ids"],
                "review_status": "New evidence-aware candidate family",
                "limitations": item["limitations"],
                "sources": unique_sources(item["sources"]),
                "stable_record": "",
                "source_access_count": len(item["sources"]),
                "source_access_fail_count": 0,
                "kind": kind,
                "severity": None,
                "publisher_fields": {},
            }
        )
    return rows


def sort_key(row: dict):
    candidates = [row.get("reference_date", ""), row.get("reported_date", "")]
    for value in candidates:
        if value and value != "Unknown":
            try:
                return (0, datetime.strptime(value[:10], "%Y-%m-%d"), row["ws2_id"])
            except ValueError:
                pass
    return (1, datetime.max, row["ws2_id"])


def source_cell(sources: list[dict], index: int) -> str:
    if index >= len(sources):
        return "No supporting source located" if index == 0 else ""
    return sources[index]["label"]


def row_values(row: dict) -> list[str]:
    additional = "\n".join(
        f"{s['label']}: {s['url']}" for s in row["sources"][1:]
    )
    return [
        row["ws2_id"],
        row["reference_date"],
        row["reported_date"],
        row["headline"],
        row["collection"],
        row["environment"],
        row["actor"],
        row["what_happened"],
        row["outcome"],
        row["real_external_action"],
        row["boundary"],
        row["evidence_status"],
        row["source_quality"],
        row["relationship"],
        row["existing_ids"],
        row["review_status"],
        row["limitations"],
        source_cell(row["sources"], 0),
        additional,
        row["stable_record"],
    ]


def add_register_sheet(wb: Workbook, title: str, rows: list[dict]) -> None:
    ws = wb.create_sheet(title)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    for row in rows:
        ws.append(row_values(row))
        excel_row = ws.max_row
        for cell in ws[excel_row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN_GREY)
        primary = row["sources"][0]["url"] if row["sources"] else None
        if primary:
            ws.cell(excel_row, 18).hyperlink = primary
            ws.cell(excel_row, 18).style = "Hyperlink"
        if row["stable_record"]:
            ws.cell(excel_row, 20).hyperlink = row["stable_record"]
            ws.cell(excel_row, 20).style = "Hyperlink"
        if row["source_access_count"] == 0:
            fill = UNREACHABLE_FILL
        elif row["review_status"].startswith("Screened"):
            fill = SCREENED_FILL
        elif row["relationship"].startswith("Existing"):
            fill = OVERLAP_FILL
        elif row["kind"] in ("demonstrated", "controlled"):
            fill = CONTROLLED_FILL
        else:
            fill = REALIZED_FILL
        ws.cell(excel_row, 1).fill = fill
        ws.row_dimensions[excel_row].height = 96
    widths = [12, 20, 22, 42, 24, 28, 25, 70, 52, 27, 28, 42, 22, 28, 17, 38, 58, 40, 58, 34]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False


def add_readme(wb: Workbook, rows: list[dict], permission: dict, reachability: dict) -> None:
    ws = wb.active
    ws.title = "Read Me"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 115
    title = ws.cell(1, 1, "AI Incident Watch — Wide Search 2")
    title.font = Font(size=18, bold=True, color="FFFFFF")
    title.fill = HEADER_FILL
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 30

    pp_count = sum(r["collection"] == "Permission Protocol" for r in rows)
    metr_count = sum(r["collection"].startswith("METR") for r in rows)
    manual_count = sum(r["collection"] == "Direct expanded search" for r in rows)
    overlaps = sum(r["relationship"].startswith("Existing") for r in rows)
    potential = sum(r["relationship"].startswith("Potential") for r in rows)
    no_source = sum(r["source_access_count"] == 0 for r in rows)
    controlled = sum(r["kind"] in ("demonstrated", "controlled") for r in rows)
    operational = sum(r["kind"] in ("realized", "operational", "authorized_game") for r in rows)
    screened = sum(r["kind"] == "screened" for r in rows)

    content = [
        ("As of", AS_OF),
        ("Purpose", "A separate, broad discovery register for possible AI-agent incidents, controlled analogues, security demonstrations, disputed reports, and useful negative cases. It does not alter the 46-family factual baseline or the dated evidence-update workbook."),
        ("Row unit", "One source record or candidate family. The 225 rows are not a unique-event census. Some tracker records describe overlapping campaigns, product-level vulnerability clusters, or different reports about the same underlying family."),
        ("Total rows", str(len(rows))),
        ("Discovery backbones", f"{pp_count} Permission Protocol records + {metr_count} METR records + {manual_count} directly researched additional families"),
        ("Broad routing", f"{operational} reported/operational/authorized-game rows; {controlled} controlled or demonstrated rows; {screened} screened METR context rows"),
        ("Reconciliation", f"{overlaps} exact existing baseline/update overlaps and {potential} possible overlap requiring adjudication. Existing IDs are listed in the register; existing factual findings control over imported tracker wording."),
        ("Source access", f"The batch check opened the first 8 KiB of {reachability['unique_urls']} unique Permission Protocol source URLs: {reachability['reachable']} were reachable and {reachability['unavailable']} were blocked, paywalled, unauthorized, or returned errors. After direct-source enrichment, {no_source} workbook rows still have no linked supporting source reachable from this environment."),
        ("Permission Protocol status", f"Imported from a public {permission['recordCount']}-record dataset updated {permission['updatedAt']}. Its own totals are {permission['statistics']['realized']} realized reports and {permission['statistics']['demonstrated']} demonstrations. Source quality, category, severity, gate, and authority fields are the publisher's assessments."),
        ("METR status", "METR's register supplies quoted passages and original-source links for 44 documented cases. Six rows without a task/permission boundary action remain in the workbook as screened context. Most other METR rows are controlled evaluations or internal deployments."),
        ("Date warning", permission["dateNote"]),
        ("Control-analysis warning", permission["controlAnalysisNote"]),
        ("Overlap warning", permission["overlapNote"]),
        ("Evidence rule", "A reachable page proves only that the source route worked. It does not prove every claim. Rows marked 'adjudication pending' require passage-level review before they are added to the master incident census. Rows with unavailable sources remain leads."),
        ("Corrections opened in this pass", "DataTalks: the founder said he watched the destructive command, did not stop it, does not blame Claude, and recovered the database. Amazon Kiro: Amazon disputes AI causation and attributes the interruption to user error and a misconfigured role. Amazon Q: the malicious code shipped but AWS says a syntax error prevented execution. EchoLeak and Gemini CLI: demonstrated vulnerabilities with no confirmed in-the-wild exploitation in the sources reviewed."),
        ("Sheet guide", "All Candidates contains every row once. Reported & Operational and Controlled & Security are views. Existing Overlaps isolates reconciled duplicates. Screened & Unreachable holds screened context and leads whose linked evidence could not be opened. Source Registry lists every clickable source. Search Log records the search routes used."),
        ("Preservation", "The three supplied datasets, 46-family master workbook, factual rebuild, and dated evidence-update workbook are unchanged."),
    ]
    for label, value in content:
        ws.append([label, value])
        r = ws.max_row
        ws.cell(r, 1).font = Font(bold=True, color="17365D")
        ws.cell(r, 1).fill = SUBHEADER_FILL
        for cell in ws[r]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN_GREY)
        ws.row_dimensions[r].height = max(30, min(120, 15 * (value.count("\n") + 2)))
    links = [
        ("Permission Protocol tracker", PERMISSION_TRACKER),
        ("Permission Protocol methodology", permission["methodologyUrl"]),
        ("METR documented incidents", METR_TRACKER),
        ("METR Frontier Risk Report", METR_REPORT),
    ]
    ws.append([])
    ws.append(["Core sources", ""])
    ws.cell(ws.max_row, 1).font = Font(bold=True, color="FFFFFF")
    ws.cell(ws.max_row, 1).fill = HEADER_FILL
    for label, url in links:
        ws.append([label, url])
        cell = ws.cell(ws.max_row, 2)
        cell.hyperlink = url
        cell.style = "Hyperlink"
    ws.freeze_panes = "A2"


def add_source_registry(wb: Workbook, rows: list[dict], reachability: dict) -> list[dict]:
    reach = {x["url"]: x for x in reachability["results"]}
    sources: dict[str, dict] = {}
    for row in rows:
        for source in row["sources"]:
            entry = sources.setdefault(
                source["url"],
                {"url": source["url"], "labels": set(), "case_ids": set(), "collections": set()},
            )
            entry["labels"].add(source["label"])
            entry["case_ids"].add(row["ws2_id"])
            entry["collections"].add(row["collection"])
    source_rows = []
    for index, url in enumerate(sorted(sources), 1):
        item = sources[url]
        check = reach.get(url)
        if url in LOCALLY_VERIFIED_URLS:
            status = "Retrieved/opened directly"
            detail = "Relevant passage reviewed in this or the baseline/update pass"
        elif check and check["status"] == "reachable":
            status = "Reachable"
            detail = f"HTTP {check.get('http_status', 'Unknown')}; first {check.get('sample_bytes', 0)} bytes opened"
        elif check:
            status = "Unavailable in environment"
            detail = check.get("error", "Unknown error")
        else:
            status = "Listed; not batch-checked"
            detail = "Source was added by the METR/manual research route"
        source_rows.append(
            {
                "source_id": f"WS2-S{index:04d}",
                "label": " | ".join(sorted(item["labels"])),
                "url": url,
                "domain": urlsplit(url).netloc,
                "status": status,
                "detail": detail,
                "case_ids": "; ".join(sorted(item["case_ids"])),
                "collections": "; ".join(sorted(item["collections"])),
            }
        )
    ws = wb.create_sheet("Source Registry")
    headers = ["Source ID", "Label", "URL", "Domain", "Access status", "Access detail", "WS2 case IDs", "Source collection(s)"]
    ws.append(headers)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for item in source_rows:
        ws.append([item["source_id"], item["label"], item["url"], item["domain"], item["status"], item["detail"], item["case_ids"], item["collections"]])
        r = ws.max_row
        ws.cell(r, 3).hyperlink = item["url"]
        ws.cell(r, 3).style = "Hyperlink"
        for c in ws[r]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=THIN_GREY)
        if item["status"] == "Unavailable in environment":
            ws.cell(r, 1).fill = UNREACHABLE_FILL
    for idx, width in enumerate([13, 48, 75, 30, 25, 50, 35, 30], 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    return source_rows


def add_search_log(wb: Workbook) -> None:
    payload = load_json(SEARCH_LOG_FILE)
    ws = wb.create_sheet("Search Log")
    headers = ["Date", "Route", "Scope / query family", "Result", "Evidence use"]
    ws.append(headers)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = Font(color="FFFFFF", bold=True)
    for item in payload["searches"]:
        ws.append([item["date"], item["route"], item["scope"], item["result"], item["evidence_use"]])
        for c in ws[ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=THIN_GREY)
    for idx, width in enumerate([14, 28, 65, 60, 52], 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False


def validate(rows: list[dict], source_rows: list[dict]) -> None:
    ids = [r["ws2_id"] for r in rows]
    assert len(ids) == len(set(ids)), "Duplicate WS2 IDs"
    assert len(rows) == 225, f"Expected 225 rows, found {len(rows)}"
    assert sum(r["collection"] == "Permission Protocol" for r in rows) == 171
    assert sum(r["collection"].startswith("METR") for r in rows) == 44
    assert sum(r["collection"] == "Direct expanded search" for r in rows) == 10
    assert all(r["sources"] for r in rows), "Every candidate needs a source or explicit no-source marker"
    source_urls = {s["url"] for s in source_rows}
    assert all(s["url"] in source_urls for r in rows for s in r["sources"])


def main() -> None:
    permission = load_json(PERMISSION_FILE)
    metr = load_json(METR_FILE)
    manual = load_json(MANUAL_FILE)
    reachability = load_json(REACHABILITY_FILE)
    assert permission["recordCount"] == 171
    assert len(permission["records"]) == 171
    assert len([x for x in metr if not x["id"].startswith("INC-H")]) == 44
    assert len(manual["cases"]) == 10

    rows = pp_rows(permission, reachability) + metr_rows(metr) + manual_rows(manual)
    rows.sort(key=sort_key)

    wb = Workbook()
    add_readme(wb, rows, permission, reachability)
    add_register_sheet(wb, "All Candidates", rows)
    reported = [r for r in rows if r["kind"] in ("realized", "operational", "authorized_game")]
    controlled = [r for r in rows if r["kind"] in ("demonstrated", "controlled")]
    overlaps = [r for r in rows if r["relationship"].startswith(("Existing", "Potential"))]
    screened = [r for r in rows if r["kind"] == "screened" or r["source_access_count"] == 0]
    add_register_sheet(wb, "Reported & Operational", reported)
    add_register_sheet(wb, "Controlled & Security", controlled)
    add_register_sheet(wb, "Existing Overlaps", overlaps)
    add_register_sheet(wb, "Screened & Unreachable", screened)
    source_rows = add_source_registry(wb, rows, reachability)
    add_search_log(wb)

    validate(rows, source_rows)
    wb.save(OUTPUT)

    # Reopen to catch malformed workbook output and validate the principal view.
    check = load_workbook(OUTPUT, read_only=False, data_only=False)
    assert check.sheetnames == [
        "Read Me", "All Candidates", "Reported & Operational", "Controlled & Security",
        "Existing Overlaps", "Screened & Unreachable", "Source Registry", "Search Log",
    ]
    assert check["All Candidates"].max_row == len(rows) + 1
    assert check["All Candidates"].freeze_panes == "A2"
    assert check["All Candidates"].auto_filter.ref
    check.close()

    result = {
        "metadata": {
            "created": AS_OF,
            "status": "discovery_register_not_confirmed_incident_census",
            "record_count": len(rows),
            "collections": dict(Counter(r["collection"] for r in rows)),
            "kinds": dict(Counter(r["kind"] for r in rows)),
            "relationships": dict(Counter(r["relationship"] for r in rows)),
            "rows_without_reachable_supporting_source": sum(r["source_access_count"] == 0 for r in rows),
            "source_registry_count": len(source_rows),
            "workbook": OUTPUT.name,
        },
        "records": rows,
        "sources": source_rows,
    }
    REGISTER_OUTPUT.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n")
    digest = hashlib.sha256(OUTPUT.read_bytes()).hexdigest()
    print(
        f"Built {OUTPUT.name}: {len(rows)} candidate/source records, "
        f"{len(reported)} reported/operational, {len(controlled)} controlled/security, "
        f"{len(overlaps)} overlaps, {len(screened)} screened/unreachable, "
        f"{len(source_rows)} unique sources; sha256={digest}"
    )


if __name__ == "__main__":
    main()
