#!/usr/bin/env python3
"""Build the dated evidence-update workbook without changing the 46-case baseline."""

from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data/monitoring/evidence_updates_2026-10-13.json"
WATCH = ROOT / "data/monitoring/case_watch_2026-10-13.json"
BASELINE = ROOT / "data/rebuilt/factual_rebuild.json"
OUTPUT = ROOT / "AI_Incident_Watch_Evidence_Updates.xlsx"

NAVY = "17365D"
BLUE = "2F75B5"
PALE_BLUE = "D9EAF7"
PALE_GREEN = "E2F0D9"
PALE_AMBER = "FFF2CC"
PALE_RED = "FCE4D6"
PALE_GREY = "E7E6E6"
WHITE = "FFFFFF"
GRID = Side(style="thin", color="C8D0D9")


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


data = load_json(DATA)
watch = load_json(WATCH)
baseline = load_json(BASELINE)
sources = {row["source_id"]: row for row in data["sources"]}


def normalize_xlsx_archive(path: Path) -> None:
    """Normalize ZIP metadata so unchanged inputs produce the same workbook bytes."""
    normalized = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        normalized, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=(2026, 10, 13, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    normalized.replace(path)


def source_label(source_id: str | None) -> str:
    return sources[source_id]["label"] if source_id else ""


def add_source_link(cell, source_id: str | None) -> None:
    if not source_id:
        cell.value = ""
        return
    source = sources[source_id]
    cell.value = source["label"]
    cell.hyperlink = source["url"]
    cell.style = "Hyperlink"
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def add_table_sheet(wb: Workbook, title: str, columns: list[tuple[str, str, int]], rows: list[dict]):
    ws = wb.create_sheet(title)
    for col_idx, (header, _, width) in enumerate(columns, 1):
        cell = ws.cell(1, col_idx, header)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 34

    for row_idx, row in enumerate(rows, 2):
        for col_idx, (_, key, _) in enumerate(columns, 1):
            cell = ws.cell(row_idx, col_idx)
            if key in {"source_1", "source_2"}:
                add_source_link(cell, row.get(key))
            else:
                value = row.get(key)
                cell.value = "Unknown" if value is None else value
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=GRID)
        ws.row_dimensions[row_idx].height = 90

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    return ws


def style_status_cells(ws, column: int) -> None:
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row, column)
        text = str(cell.value or "").lower()
        if "supported" in text or "corroborated" in text:
            color = PALE_GREEN
        elif "unverified" in text or "pending" in text or "qualified" in text:
            color = PALE_AMBER
        elif "unsupported" in text or "screened" in text:
            color = PALE_RED
        else:
            color = PALE_GREY
        cell.fill = PatternFill("solid", fgColor=color)


def make_dashboard(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Read Me"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 27
    ws.column_dimensions["B"].width = 105
    ws.merge_cells("A1:B1")
    ws["A1"] = "AI Incident Watch — Evidence Update Register"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=16)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34

    rows = [
        ("Baseline", "The 46-family master reading list remains the factual baseline as of 9 October 2026. This workbook does not silently rewrite it."),
        ("This scan", "Sources published or posted through 13 October 2026, plus older sources found during the scan."),
        ("New since baseline", f"{len(data['new_cases'])} provisional families. These are candidates for adjudication, not automatically confirmed incidents."),
        ("Existing case updates", f"{len(data['existing_updates'])} evidence updates affecting baseline cases. One strengthens R020; one adds cross-case Anthropic remediation."),
        ("Backfill found now", f"{len(data['backfills'])} pre-baseline sources were missing from the corpus and need adjudication."),
        ("46 Case Watch", "Every baseline family appears exactly once with its targeted query, candidate-headline count, and material-change decision."),
        ("Evidence rule", data["metadata"]["decision_rule"]),
        ("Allegations", "A direct user post proves that the allegation was made. It does not by itself prove product causation, command execution, damage scale, or recovery outcome."),
        ("Date rule", "Unknown means the source did not disclose a reliable date. The report/post date is kept separately and is never substituted silently for an incident date."),
        ("How to update", "Run scripts/scan_evidence_radar.py for discovery, adjudicate candidate sources, update the dated JSON files, then run scripts/build_evidence_updates.py."),
        ("Notification limit", "This chat cannot initiate a message while no session is active. The saved last-checked date makes the next scan incremental. External push notification requires a scheduled channel such as GitHub Issues or email."),
    ]
    for idx, (label, value) in enumerate(rows, 3):
        ws.cell(idx, 1, label).font = Font(bold=True, color=NAVY)
        ws.cell(idx, 1).fill = PatternFill("solid", fgColor=PALE_BLUE)
        ws.cell(idx, 2, value)
        for col in range(1, 3):
            ws.cell(idx, col).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(idx, col).border = Border(bottom=GRID)
        ws.row_dimensions[idx].height = 52

    ws["A16"] = "Workbook tabs"
    ws["A16"].font = Font(bold=True, color=WHITE)
    ws["A16"].fill = PatternFill("solid", fgColor=BLUE)
    ws["B16"].fill = PatternFill("solid", fgColor=BLUE)
    tabs = [
        ("Existing Case Updates", "Evidence that changes or extends a baseline case."),
        ("New Since Baseline", "Provisional new case families reported on or after the baseline day."),
        ("Backfill Found Now", "Older evidence discovered during this scan."),
        ("Screened Candidates", "Items reviewed but not promoted, with the reason."),
        ("46 Case Watch", "One-row-per-baseline-family scan log."),
        ("Source Registry", "Web links, source type, local preservation path, size, and SHA-256."),
        ("Run Log", "Checkpoint for the next incremental scan."),
    ]
    for idx, (label, value) in enumerate(tabs, 17):
        ws.cell(idx, 1, label).font = Font(bold=True)
        ws.cell(idx, 2, value)
        for col in range(1, 3):
            ws.cell(idx, col).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(idx, col).border = Border(bottom=GRID)
        ws.row_dimensions[idx].height = 32
    ws.freeze_panes = "A3"


def build() -> None:
    baseline_ids = [case["canonical_id"] for case in baseline["cases"]]
    watch_ids = [case["case_id"] for case in watch["cases"]]
    if len(baseline_ids) != 46 or len(set(baseline_ids)) != 46:
        raise ValueError("Baseline must contain 46 unique case families")
    if watch_ids != baseline_ids:
        raise ValueError("46 Case Watch must include every baseline family once and in baseline order")
    if len({row["radar_id"] for row in data["new_cases"]}) != len(data["new_cases"]):
        raise ValueError("New provisional family IDs must be unique")

    wb = Workbook()
    make_dashboard(wb)

    ws = add_table_sheet(wb, "Existing Case Updates", [
        ("Update ID", "update_id", 12), ("First seen", "first_seen", 13),
        ("Published", "published", 24), ("Baseline case(s)", "baseline_cases", 18),
        ("Update type", "update_type", 22), ("New information", "new_information", 58),
        ("Effect on baseline finding", "effect_on_finding", 52), ("Proposed action", "proposed_action", 48),
        ("Source 1", "source_1", 35), ("Source 2", "source_2", 35), ("Limitations", "limitations", 48),
    ], data["existing_updates"])
    style_status_cells(ws, 5)

    ws = add_table_sheet(wb, "New Since Baseline", [
        ("Radar ID", "radar_id", 12), ("First seen", "first_seen", 13),
        ("Date reported", "reported", 23), ("Incident date", "incident_date", 30),
        ("Provisional family", "family", 36), ("Relationship", "relationship", 24),
        ("Environment", "environment", 34), ("What happened", "what_happened", 66),
        ("Outcome", "outcome", 44), ("Evidence status", "evidence_status", 40),
        ("Materiality", "materiality", 13), ("Proposed corpus action", "proposed_action", 48),
        ("Source 1", "source_1", 35), ("Source 2", "source_2", 35), ("Limitations / correction", "limitations", 52),
    ], data["new_cases"])
    style_status_cells(ws, 10)

    ws = add_table_sheet(wb, "Backfill Found Now", [
        ("Backfill ID", "backfill_id", 13), ("First seen", "first_seen", 13),
        ("Published", "published", 14), ("Incident date", "incident_date", 27),
        ("Candidate family", "family", 39), ("Relationship", "relationship", 42),
        ("What happened", "what_happened", 66), ("Outcome", "outcome", 48),
        ("Evidence status", "evidence_status", 42), ("Proposed corpus action", "proposed_action", 52),
        ("Source 1", "source_1", 38), ("Source 2", "source_2", 34), ("Limitations", "limitations", 50),
    ], data["backfills"])
    style_status_cells(ws, 9)

    ws = add_table_sheet(wb, "Screened Candidates", [
        ("Screen ID", "screen_id", 12), ("Published", "published", 22),
        ("Candidate", "candidate", 42), ("Decision", "decision", 44),
        ("Reason", "reason", 70), ("Source 1", "source_1", 38), ("Source 2", "source_2", 34),
    ], data["screened"])
    style_status_cells(ws, 4)

    ws = add_table_sheet(wb, "46 Case Watch", [
        ("Case ID", "case_id", 11), ("Baseline family", "title", 42),
        ("Baseline claim status", "baseline_claim_status", 22), ("Baseline disposition", "baseline_disposition", 28),
        ("Targeted query", "query", 55), ("Headline candidates", "headline_candidates", 18),
        ("Scan result", "scan_result", 28), ("Review note", "review_note", 60), ("Last checked", "last_checked", 14),
    ], watch["cases"])
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 54
    style_status_cells(ws, 7)

    source_rows = []
    for source in data["sources"]:
        local_path = ROOT / source["local_file"]
        if not local_path.exists():
            raise FileNotFoundError(local_path)
        body = local_path.read_bytes()
        source_rows.append({
            **source,
            "web_link": source["label"],
            "bytes": len(body),
            "sha256": hashlib.sha256(body).hexdigest(),
        })
    ws = add_table_sheet(wb, "Source Registry", [
        ("Source ID", "source_id", 12), ("Source", "web_link", 45),
        ("Published", "published", 16), ("Evidence type", "type", 42),
        ("Local preserved file", "local_file", 58), ("Bytes", "bytes", 14),
        ("SHA-256", "sha256", 68), ("Use / limitations", "notes", 72),
    ], source_rows)
    for row_idx, source in enumerate(source_rows, 2):
        cell = ws.cell(row_idx, 2)
        cell.hyperlink = source["url"]
        cell.style = "Hyperlink"
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row_idx].height = 72

    add_table_sheet(wb, "Run Log", [
        ("Run date", "run_date", 15), ("Baseline cutoff", "baseline_cutoff", 17),
        ("Method", "method", 88), ("Existing-case material updates", "existing_updates", 26),
        ("New provisional families", "new_families", 24), ("Backfills", "backfills", 14),
        ("Screened candidates", "screened", 21), ("Next scan starts after", "next_after", 22),
    ], [{
        "run_date": data["metadata"]["scan_as_of"],
        "baseline_cutoff": data["metadata"]["baseline_as_of"],
        "method": data["metadata"]["method"],
        "existing_updates": len(data["existing_updates"]),
        "new_families": len(data["new_cases"]),
        "backfills": len(data["backfills"]),
        "screened": len(data["screened"]),
        "next_after": data["metadata"]["scan_as_of"],
    }])

    wb.properties.title = "AI Incident Watch Evidence Updates"
    wb.properties.subject = "Delta evidence register after the 46-case baseline"
    wb.properties.creator = "AI Incident Watch factual rebuild"
    fixed_time = datetime(2026, 10, 13, tzinfo=timezone.utc)
    wb.properties.created = fixed_time
    wb.properties.modified = fixed_time
    wb.save(OUTPUT)
    normalize_xlsx_archive(OUTPUT)

    # Re-open for a structural validation independent of in-memory state.
    check = load_workbook(OUTPUT, read_only=False, data_only=False)
    expected = [
        "Read Me", "Existing Case Updates", "New Since Baseline", "Backfill Found Now",
        "Screened Candidates", "46 Case Watch", "Source Registry", "Run Log",
    ]
    if check.sheetnames != expected:
        raise ValueError(f"Unexpected sheet set: {check.sheetnames}")
    ids = [check["46 Case Watch"].cell(row, 1).value for row in range(2, check["46 Case Watch"].max_row + 1)]
    if ids != baseline_ids:
        raise ValueError("Saved workbook does not contain the 46 baseline IDs exactly once")
    if check["New Since Baseline"].max_row - 1 != len(data["new_cases"]):
        raise ValueError("Saved workbook lost provisional families")
    if not any(check["New Since Baseline"].cell(2, col).hyperlink for col in range(1, check["New Since Baseline"].max_column + 1)):
        raise ValueError("Saved workbook has no source hyperlink in first provisional row")
    check.close()
    print(
        f"Built {OUTPUT.name}: 46 watched baseline families, "
        f"{len(data['new_cases'])} provisional new families, "
        f"{len(data['backfills'])} backfills, {len(data['sources'])} preserved sources"
    )


if __name__ == "__main__":
    build()
