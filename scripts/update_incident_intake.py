#!/usr/bin/env python3
"""Upsert GitHub incident issue to an UNVERIFIED intake register only."""
import json, os, re
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE=Path(__file__).resolve().parents[1]/"data"/"intake"
J=BASE/"incident_leads.json"; X=BASE/"incident_leads.xlsx"
COLUMNS=[("issue_number","Issue #"),("issue_url","Issue link"),("event","Event or finding"),("source_url","Source"),("occurrence_date","Occurrence date"),("source_context","Context"),("what_happened","What happened"),("notes","Notes / uncertainty"),("evidence_status","Evidence status"),("updated_at","Updated")]

def fields(body):
    parts=re.split(r"(?m)^### (.+?)\s*$",body or "")
    return {parts[i].strip():parts[i+1].strip() for i in range(1,len(parts)-1,2)}

def csv_safe(value):
    if isinstance(value,str) and value.startswith(("=","+","-","@")): return "'"+value
    return value

def main():
    BASE.mkdir(parents=True,exist_ok=True)
    rows=json.loads(J.read_text()) if J.exists() else []
    if os.environ.get("GITHUB_EVENT_NAME") == "issues":
        event=json.loads(Path(os.environ["GITHUB_EVENT_PATH"]).read_text())
        issue=event["issue"]; owner=os.environ["GITHUB_REPOSITORY"].split("/")[0]
        if issue["user"]["login"] != owner or not issue["title"].startswith("[Incident]"):
            raise ValueError("Only owner-authored incident issues are accepted")
        f=fields(issue.get("body",""))
        required=("Event or finding","Source URL","What happened")
        if any(f.get(k,"") in ("","_No response_") for k in required):
            raise ValueError("Required form fields are missing")
        url=f["Source URL"].strip().split()[0].strip("<>")
        parsed=urlparse(url)
        if parsed.scheme not in ("http","https") or not parsed.netloc:
            raise ValueError("A valid source URL is required")
        row={"issue_number":issue["number"],"issue_url":issue["html_url"],
             "event":f["Event or finding"],"source_url":url,
             "occurrence_date":f.get("Occurrence date",""),
             "source_context":f.get("Source context","Uncertain"),
             "what_happened":f["What happened"],
             "notes":f.get("Evidence gaps or possible duplicates",""),
             "evidence_status":"Unverified lead",
             "updated_at":issue.get("updated_at","")}
        rows=[r for r in rows if r["issue_number"]!=row["issue_number"]]+[row]
    rows.sort(key=lambda r:r["issue_number"])
    J.write_text(json.dumps(rows,ensure_ascii=False,indent=2)+"\n")
    wb=Workbook();ws=wb.active;ws.title="Unverified intake";ws.append([label for _,label in COLUMNS])
    ws.freeze_panes="C2";ws.auto_filter.ref=f"A1:J{len(rows)+1}"
    for i,width in enumerate([12,42,42,50,20,35,85,70,22,26],1):
        ws.column_dimensions[get_column_letter(i)].width=width
        c=ws.cell(1,i);c.fill=PatternFill("solid",fgColor="17365D");c.font=Font(color="FFFFFF",bold=True)
    for r in rows:
        ws.append([csv_safe(r.get(key,"")) for key,_ in COLUMNS])
        for cell in ws[ws.max_row]:cell.alignment=Alignment(wrap_text=True,vertical="top")
        for idx in (2,4):
            c=ws.cell(ws.max_row,idx)
            if str(c.value).startswith(("https://","http://")):
                c.hyperlink=c.value;c.style="Hyperlink"
    wb.save(X)

if __name__=="__main__":main()
